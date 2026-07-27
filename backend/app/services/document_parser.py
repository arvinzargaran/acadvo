import base64
import csv
import io
import json

from app.ai.client import get_groq_client
from app.config import settings

SPREADSHEET_MIME_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "text/csv",
}

EXTRACTION_PROMPT = """Analyze this degree document (academic calendar, degree checklist, transcript, or similar).
Extract ALL degree requirements and courses into this exact JSON format:

{
  "requirements": [
    {"bucket_name": "...", "credits_required": 0.0, "description": "..."}
  ],
  "courses": [
    {
      "course_code": "...",
      "course_name": "...",
      "credit_value": 0.0,
      "requirement_bucket": "...",
      "term": null,
      "status": "completed",
      "grade": null
    }
  ]
}

Rules:
- Group courses into logical requirement buckets (e.g. "Major Requirements", "Science Breadth", "Electives")
- Match each course to the most appropriate requirement bucket
- status MUST be one of: "completed", "in_progress", "planned", "not_scheduled"
- If a course has a grade, mark status as "completed"
- If a course appears to be currently enrolled, mark status as "in_progress"
- If no clear status, mark as "planned"
- If no term is visible, set term to null
- If no grade is visible, set grade to null
- credit_value should be numeric (e.g. 0.5, 1.0, 3.0)
- credits_required per bucket should be the total credits needed for that bucket
- Return ONLY valid JSON, no markdown fences, no extra text"""


def _spreadsheet_to_text(file_bytes: bytes, mime_type: str) -> str:
    if mime_type == "text/csv":
        return file_bytes.decode("utf-8", errors="replace")

    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    lines = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        lines.append(f"=== Sheet: {sheet_name} ===")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) if c is not None else "" for c in row]
            lines.append("\t".join(cells))
    wb.close()
    return "\n".join(lines)


def _pdf_to_text(file_bytes: bytes) -> str:
    import pdfplumber
    text_parts = []
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                text_parts.append(text)
    return "\n\n".join(text_parts)


async def parse_degree_document(file_bytes: bytes, mime_type: str) -> dict:
    client = get_groq_client()

    if mime_type == "application/pdf":
        content_text = _pdf_to_text(file_bytes)
        response = client.chat.completions.create(
            model=settings.GROQ_MODEL,
            messages=[{"role": "user", "content": f"{EXTRACTION_PROMPT}\n\nDocument content:\n\n{content_text}"}],
            max_tokens=4096,
        )
    elif mime_type in SPREADSHEET_MIME_TYPES:
        spreadsheet_text = _spreadsheet_to_text(file_bytes, mime_type)
        response = client.chat.completions.create(
            model=settings.GROQ_MODEL,
            messages=[{"role": "user", "content": f"{EXTRACTION_PROMPT}\n\nSpreadsheet content:\n\n{spreadsheet_text}"}],
            max_tokens=4096,
        )
    else:
        # Image — use vision model
        b64 = base64.b64encode(file_bytes).decode()
        response = client.chat.completions.create(
            model=settings.GROQ_VISION_MODEL,
            messages=[{
                "role": "user",
                "content": [
                    {"type": "text", "text": EXTRACTION_PROMPT},
                    {"type": "image_url", "image_url": {"url": f"data:{mime_type};base64,{b64}"}},
                ],
            }],
            max_tokens=4096,
        )

    text = response.choices[0].message.content.strip()

    if text.startswith("```"):
        text = text.split("\n", 1)[1].rsplit("```", 1)[0].strip()

    try:
        parsed = json.loads(text)
    except json.JSONDecodeError:
        raise ValueError(
            "The AI response was not valid JSON. "
            "Try uploading a clearer document or a different format."
        )

    for req in parsed.get("requirements", []):
        req["bucket_name"] = req.get("bucket_name") or ""
        try:
            req["credits_required"] = float(req.get("credits_required") or 0)
        except (TypeError, ValueError):
            req["credits_required"] = 0.0

    _VALID_STATUSES = {"completed", "in_progress", "planned", "not_scheduled"}
    _STATUS_MAP = {
        "not completed": "planned",
        "incomplete": "in_progress",
        "enrolled": "in_progress",
        "registered": "in_progress",
        "current": "in_progress",
        "taking": "in_progress",
        "done": "completed",
        "passed": "completed",
        "failed": "completed",
        "withdrawn": "not_scheduled",
        "dropped": "not_scheduled",
    }

    for course in parsed.get("courses", []):
        course["course_code"] = course.get("course_code") or ""
        course["course_name"] = course.get("course_name") or ""
        try:
            course["credit_value"] = float(course.get("credit_value") or 0)
        except (TypeError, ValueError):
            course["credit_value"] = 0.0
        raw_status = (course.get("status") or "not_scheduled").strip().lower()
        if raw_status in _VALID_STATUSES:
            course["status"] = raw_status
        else:
            course["status"] = _STATUS_MAP.get(raw_status, "planned")

    return parsed
