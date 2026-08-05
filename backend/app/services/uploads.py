"""Validation and bounds for user-supplied file uploads.

Two problems the routers had before this module existed:

1. **The declared type was trusted.** ``UploadFile.content_type`` is just a
   header the client writes; checking it against an allowlist proves nothing.
   A file labelled ``image/png`` was handed straight to the vision model, and
   one labelled ``application/pdf`` straight to pdfplumber, whatever its actual
   bytes were. Sniffing the magic number is what actually establishes the type.

2. **Parsing was unbounded.** A 10 MB upload is not a 10 MB workload: XLSX and
   PDF are both compressed container formats, so a small file can expand into
   millions of cells or tens of thousands of pages. Every extraction here is now
   capped on the dimensions that actually blow up, and stops early rather than
   building an unbounded string in memory.
"""

import io

from fastapi import HTTPException, status

# ── Type sniffing ────────────────────────────────────────────────────

PDF = "application/pdf"
PNG = "image/png"
JPEG = "image/jpeg"
WEBP = "image/webp"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
XLS = "application/vnd.ms-excel"
CSV = "text/csv"

# Client-declared types we accept, mapped to the types the bytes are allowed to
# actually be. `image/jpg` is not a real MIME type but browsers still send it.
_DECLARED_TO_ALLOWED_SNIFFED: dict[str, set[str]] = {
    PDF: {PDF},
    PNG: {PNG},
    JPEG: {JPEG},
    "image/jpg": {JPEG},
    WEBP: {WEBP},
    # Both spreadsheet types are ZIP containers or OLE2; CSV is plain text and
    # cannot be sniffed, so it is handled separately below.
    XLSX: {XLSX, XLS},
    XLS: {XLSX, XLS},
    CSV: {CSV},
}


def sniff_mime_type(data: bytes) -> str | None:
    """Identify a file from its magic bytes. Returns None if unrecognised."""
    if data.startswith(b"%PDF-"):
        return PDF
    if data.startswith(b"\x89PNG\r\n\x1a\n"):
        return PNG
    if data.startswith(b"\xff\xd8\xff"):
        return JPEG
    if data[:4] == b"RIFF" and data[8:12] == b"WEBP":
        return WEBP
    # XLSX is a ZIP archive; legacy XLS is an OLE2 compound file.
    if data.startswith(b"PK\x03\x04") or data.startswith(b"PK\x05\x06"):
        return XLSX
    if data.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return XLS
    return None


def validate_upload(data: bytes, declared_type: str | None, allowed: set[str]) -> str:
    """Check an upload's declared type against the allowlist AND its real bytes.

    Returns the effective content type to parse with — the sniffed one, so
    downstream code branches on what the file *is*, not on what it claims.
    """
    if declared_type not in allowed:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Unsupported file type: {declared_type}.",
        )
    if not data:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST, detail="File is empty."
        )

    sniffed = sniff_mime_type(data)

    # CSV has no magic number. Accept it only if it decodes as text and does not
    # sniff as some *other* binary format wearing a .csv label.
    if declared_type == CSV:
        if sniffed is not None:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="File contents do not match the declared type.",
            )
        try:
            data[:4096].decode("utf-8")
        except UnicodeDecodeError:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail="CSV file must be valid UTF-8 text.",
            )
        return CSV

    if sniffed is None or sniffed not in _DECLARED_TO_ALLOWED_SNIFFED.get(
        declared_type, set()
    ):
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="File contents do not match the declared type.",
        )
    return sniffed


# ── Bounded extraction ───────────────────────────────────────────────
#
# Limits are generous for the real use case (a degree checklist, a transcript,
# a course syllabus, a resume) and far below what it takes to exhaust a worker.

MAX_PDF_PAGES = 100
MAX_SPREADSHEET_SHEETS = 20
MAX_SPREADSHEET_ROWS = 5_000
MAX_EXTRACTED_CHARS = 400_000  # ~100k tokens; well past any real document


class UploadTooComplexError(Exception):
    """Raised when a file is within the byte limit but explodes on parse."""


def _truncate(parts: list[str]) -> str:
    text = "\n\n".join(parts)
    if len(text) > MAX_EXTRACTED_CHARS:
        return text[:MAX_EXTRACTED_CHARS] + "\n\n[document truncated]"
    return text


def pdf_to_text(file_bytes: bytes) -> str:
    """Extract text from a PDF, bounded by page count and total characters.

    Blocking and CPU-bound — call it through ``run_blocking``, never directly
    from an async handler.
    """
    import pdfplumber

    parts: list[str] = []
    total = 0
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        if len(pdf.pages) > MAX_PDF_PAGES:
            raise UploadTooComplexError(
                f"PDF has more than {MAX_PDF_PAGES} pages."
            )
        for page in pdf.pages:
            text = page.extract_text()
            if not text:
                continue
            parts.append(text)
            total += len(text)
            # Stop reading rather than building a huge string we would only
            # truncate afterwards.
            if total > MAX_EXTRACTED_CHARS:
                break
    return _truncate(parts)


def spreadsheet_to_text(file_bytes: bytes, mime_type: str) -> str:
    """Extract text from a spreadsheet, bounded by sheets, rows and characters.

    Blocking and CPU-bound — call it through ``run_blocking``.
    """
    if mime_type == CSV:
        text = file_bytes.decode("utf-8", errors="replace")
        return text[:MAX_EXTRACTED_CHARS]

    import openpyxl

    # read_only streams rows instead of materialising the workbook, which is what
    # keeps a decompression bomb from being an instant OOM.
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    try:
        parts: list[str] = []
        total = 0
        for sheet_name in wb.sheetnames[:MAX_SPREADSHEET_SHEETS]:
            ws = wb[sheet_name]
            parts.append(f"=== Sheet: {sheet_name} ===")
            for row_index, row in enumerate(ws.iter_rows(values_only=True)):
                if row_index >= MAX_SPREADSHEET_ROWS:
                    parts.append("[sheet truncated]")
                    break
                line = "\t".join(str(c) if c is not None else "" for c in row)
                parts.append(line)
                total += len(line)
                if total > MAX_EXTRACTED_CHARS:
                    break
            if total > MAX_EXTRACTED_CHARS:
                break
    finally:
        wb.close()
    return _truncate(parts)
